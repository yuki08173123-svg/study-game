# 英検の単熟語アプリの確認用 Excel を作る（2級・3級の assemble.py から呼ぶ）
# 参考の表（英検3級_単語1400_熟語400.xlsx）にならって、
# 「使い方」シート（覚えた数を自動で数える）・学習状況の選択欄・「出典と注意」シートをつける。
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

STATUS = '"未着手,復習中,覚えた"'


def make_xlsx(out, grade, color, url, app_url, W, I, check_note):
    """grade='2級' など。W/I は単語・熟語の dict のリスト。check_note は過去問での確認の一文。"""
    wb = Workbook()
    hdr = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor=color)
    title = Font(bold=True, size=16)
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical='top')

    # 使い方
    ws = wb.active
    ws.title = '使い方'
    ws.append([f'英検{grade} 単熟語リスト'])
    ws['A1'].font = title
    ws.append([f'英単語{len(W):,}語・英熟語{len(I)}個（アプリ「英検{grade} 単熟語」と同じ並び・同じ中身）'])
    ws.append(['公式の指定語彙表ではありません。例文と覚え方はすべてオリジナルです。'])
    ws.append([])
    rows = [
        ('アプリ', app_url),
        ('収録語数', len(W)),
        ('収録熟語数', len(I)),
        ('覚えた単語', f'=COUNTIF(\'英単語{len(W)}\'!G2:G{len(W) + 1},"覚えた")'),
        ('覚えた熟語', f'=COUNTIF(\'英熟語{len(I)}\'!E2:E{len(I) + 1},"覚えた")'),
        ('復習中', f'=COUNTIF(\'英単語{len(W)}\'!G2:G{len(W) + 1},"復習中")+COUNTIF(\'英熟語{len(I)}\'!E2:E{len(I) + 1},"復習中")'),
        ('計画', '単語は1日200語×7日で1周、熟語は1日100個×4日で1周。8周くり返す（単語→熟語の順）'),
        ('学習の順番', '①英語を見て意味を言う ②意味を見て英語を言う ③例文の[ ]を隠して言う'),
        ('復習の目安', 'まちがえたものは 翌日・3日後・1週間後 にもう一度（アプリのリマインドテストと同じ）'),
        ('仕上げ', '公式過去問で 読解・リスニング・英作文・面接 を練習する'),
        ('学習チェック', '各一覧の「学習状況」で 未着手／復習中／覚えた を選ぶと、上の数が自動で変わる'),
    ]
    for k, v in rows:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 90

    def sheet(name, cols, data, widths, status_col):
        s = wb.create_sheet(name)
        s.append(cols)
        for c in s[1]:
            c.font = hdr
            c.fill = fill
        for r in data:
            s.append(r)
        for i, w in enumerate(widths):
            s.column_dimensions[chr(65 + i)].width = w
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        s.freeze_panes = 'D2'
        rng = f'{status_col}2:{status_col}{len(data) + 1}'
        dv = DataValidation(type='list', formula1=STATUS, allow_blank=True)
        dv.add(rng)
        s.add_data_validation(dv)
        s.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"覚えた"'], fill=PatternFill('solid', fgColor='D7EEDF')))
        s.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"復習中"'], fill=PatternFill('solid', fgColor='FCEFC7')))
        s.auto_filter.ref = f'A1:{chr(64 + len(cols))}{len(data) + 1}'

    sheet(f'英単語{len(W)}', ['No.', 'Day', '単語', '品詞', '発音', '意味', '学習状況', 'ほかの意味', '派生語', '例文', '訳', '覚え方'],
          [[r['n'], (r['n'] - 1) // 200 + 1, r['w'], r['p'], r.get('ph', ''), r['m'], None, ' / '.join(r.get('s', [])),
            ' / '.join(f"{x['w']}（{x['p']}）{x['m']}" for x in r.get('d', [])), r['ex'], r['ft'], r['tip']] for r in W],
          [6, 5, 16, 7, 16, 18, 10, 24, 30, 46, 40, 60], 'G')
    sheet(f'英熟語{len(I)}', ['No.', 'Day', '熟語', '意味', '学習状況', 'ほかの意味', '例文', '訳', '覚え方'],
          [[r['n'], (r['n'] - 1) // 100 + 1, r['w'], r['m'], None, ' / '.join(r.get('s', [])), r['ex'], r['ft'], r['tip']] for r in I],
          [6, 5, 26, 22, 10, 24, 46, 40, 60], 'E')

    # 出典と注意
    ws = wb.create_sheet('出典と注意')
    ws.append(['出典と注意'])
    ws['A1'].font = title
    ws.append([])
    ws.append(['資料', '使い方', 'URL'])
    for c in ws[3]:
        c.font = hdr
        c.fill = fill
    for r in [
        (f'英検{grade} 試験内容・過去問', '級の水準・題材の確認と、語彙問題での照合', url),
        ('英検 各級の審査基準', '級の水準の確認', 'https://www.eiken.or.jp/eiken/exam/criteria/index.html'),
        ('wordfreq（語の使われる頻度）', '見出し語を選ぶときの目安', 'https://github.com/rspeer/wordfreq'),
    ]:
        ws.append(list(r))
    ws.append([])
    ws.append(['注意'])
    ws.cell(ws.max_row, 1).font = bold
    for t in [
        'この一覧は学習用に編集したもので、英検協会の公式単語リストではありません。',
        check_note,
        f'この{len(W) + len(I):,}項目だけで合格を保証するものではありません。英作文・リスニング・面接も公式過去問で練習してください。',
        '多義語は、試験でよく問われる意味を「意味」に、ほかの意味を「ほかの意味」に短く書いています。文脈で意味が変わります。',
        '例文・訳・覚え方はすべてオリジナルです（市販の単語帳からは写していません）。',
    ]:
        ws.append([t])
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 44
    ws.column_dimensions['C'].width = 70
    wb.save(out)
    return out
